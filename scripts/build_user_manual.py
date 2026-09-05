#!/usr/bin/env python3
"""Builds the bilingual (English + French) user manual as Markdown and a
green-branded .docx from one source.

The document is one manual in two languages: the English version first, then a
page break, then the full French version. Every task is written as numbered
steps. The PDF is produced from the .docx by scripts/convert-docx-pdf.ps1 (Word).

    python scripts/build_user_manual.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(ROOT, 'public', 'logo.png')

GREEN_DARK = RGBColor(0x14, 0x53, 0x2D)
GREEN = RGBColor(0x15, 0x80, 0x3D)
GREEN_HEX = '15803D'
GREY = RGBColor(0x44, 0x44, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SCHOOL = os.environ.get('NEXT_PUBLIC_SCHOOL_NAME', 'Fondation Révélation Sainte Thérèse')

def h1(t): return ('h1', t)
def h2(t): return ('h2', t)
def p(t): return ('p', t)
def bullets(items): return ('bullets', items)
def steps(items): return ('steps', items)
def table(headers, rows): return ('table', headers, rows)
def note(t): return ('note', t)
PB = ('pagebreak',)

# ============================================================ ENGLISH
EN = [
    h1('About this manual'),
    p('This manual explains what the School Uniform Sales & Stock system does and, '
      'step by step, how each user performs their tasks. It covers Phase 1 — '
      'uniform production, sales, orders, alterations, returns and reporting. Read '
      'the section for your role plus the shared sections (Getting started, Key '
      'rules, Reporting a problem). The French version follows the English one.'),
    note('The interface is bilingual. Switch language any time from the top bar; '
         'every screen and receipt is available in both. What you type — names, '
         'notes, reasons — is kept exactly as written and never translated.'),

    h1('The system at a glance'),
    p('Uniforms are made and sold in the same place. The system records production '
      'as garments are finished, records sales and prints receipts, takes orders '
      'for garments not yet made, tracks alterations, handles returns and '
      'exchanges, and produces the end-of-day cash reconciliation.'),
    h2('Rules everyone must know'),
    bullets([
        'Nothing is ever edited or deleted — a mistake is corrected with a new, '
        'linked entry (a cancellation, a return); both stay visible.',
        'Everything is audited: who did what, and when.',
        'Payment is in full, on the spot — cash or mobile money, no instalments.',
        'Size is chosen at the moment of sale/order/exchange; stock is counted per '
        'size.',
        'Mobile-first: large buttons, searchable lists, minimal typing.',
    ]),

    PB,
    h1('Getting started (all users)'),
    h2('Sign in'),
    steps([
        'Open the app in a browser. There is no public sign-up — the Super Admin '
        'creates your account.',
        'Enter your email and password, then select Login.',
        'On your first login, set a new password (at least 8 characters, with a '
        'letter and a number) before continuing.',
    ]),
    p('If your account is deactivated you cannot sign in — contact your '
      'administrator. Forgotten passwords are reset by the Super Admin (no email '
      'self-service). You are signed out after inactivity: 12 hours for the Seller, '
      '2 hours for everyone else.'),
    h2('Find your way around'),
    steps([
        'Use the sidebar to move between screens (grouped Overview, Operations, '
        'Records, Administration) — you only see what your role allows.',
        'Use the search box to find a transaction by reference number, student '
        'name, parent name or phone.',
        'Open the bell for notifications; use the language and theme toggles in the '
        'top bar.',
    ]),

    PB,
    h1('Roles at a glance'),
    table(['Role', 'Who', 'What they do'], [
        ['Seller', 'Mr. Ateba (1)',
         'Production, sales, orders, alterations, returns, exchanges, '
         'cancellations. Cannot change prices.'],
        ['Administration', 'Founder, admin staff (~5)',
         'Read-only everywhere: views, reports, exports, audit log. No writing.'],
        ['Maintenance', 'Developers (2)', 'Full functional access, fully audited.'],
        ['Super Admin', '1', 'Accounts, catalogue and prices, size set, return '
         'policy windows.'],
    ]),

    PB,
    h1('The Seller'),
    p('The Seller (Mr. Ateba) runs the shop. The landing screen is Open jobs.'),
    h2('Read the Open jobs list'),
    steps([
        'Open the app — Open jobs is your home screen.',
        'Read each card: student and class, garment and size, status, date placed, '
        'and days open. Oldest is at the top.',
        'Tap a card to change its status in one tap, or to open the job.',
    ]),
    h2('Record a sale'),
    steps([
        'Open Sales.',
        'Type the student name and class, and the parent name (phone optional).',
        'Tap Add item, choose the product, then tap the size on the size bar (or '
        'type a custom size), and set the quantity. Repeat for more items.',
        'To reduce a price, enter an amount in Discount and a reason (the unit '
        'price itself cannot be edited).',
        'Choose the payment method — Cash, MoMo or Orange Money — and, for mobile '
        'money, the transaction reference.',
        'For cash, optionally enter the amount tendered to see the change.',
        'Select Save. The receipt opens and stock for that product/size goes down.',
    ]),
    note('Selling more than the system shows warns you but does not block the sale; '
         'confirm to proceed and the override is recorded.'),
    h2('Print or reprint a receipt'),
    steps([
        'After saving, the receipt opens — choose A5 or A4 and select Print.',
        'To reprint later, open the sale (via search or the sales list) and select '
        'Reprint. The copy is stamped DUPLICATA / DUPLICATE and is recorded.',
    ]),
    h2('Take an order (garment not available yet)'),
    steps([
        'Open Orders and enter the customer and line(s) with size, as for a sale.',
        'Take payment in full — the receipt is marked COMMANDE / ORDER.',
        'The order appears in Open jobs; advance it Ordered → In production → Ready '
        'as the garment is made.',
        'When the parent comes, open the order and record the collection: a '
        'collection slip prints (showing the COL and the ORD it closes) and stock '
        'is reduced at that point.',
    ]),
    h2('Take in an alteration'),
    steps([
        'Open Alterations → New.',
        'Record the parent/student, the garment, and what must be done (required).',
        'Set whether there is a charge; if so, take payment and method.',
        'Give the parent the printed deposit slip.',
        'Advance it Received → In progress → Ready, and record Returned when the '
        'parent collects. Alterations never touch stock.',
    ]),
    h2('Handle a return or exchange'),
    steps([
        'Open the original sale and select Start return (or Returns → New).',
        'Read the on-screen verdict — how long ago the sale was and whether an '
        'exchange/refund is within policy.',
        'Choose the garment(s) coming back and declare the condition (Unworn or '
        'Worn).',
        'For an exchange, pick the replacement product and size; collect the '
        'difference if dearer, refund it if cheaper. For a return, issue the refund '
        'and record its method.',
        'Enter a reason (required). If outside policy, add the reason it asks for '
        'and proceed — the override is flagged in reports.',
        'Select Save — a return/exchange receipt prints and stock updates both ways.',
    ]),
    h2('Cancel a sale'),
    steps([
        'Open the sale and select Cancel.',
        'Enter a reason (required) and confirm.',
        'Stock is put back; the sale stays visible marked Cancelled, keeps its '
        'reference, and is excluded from revenue.',
    ]),
    h2('Record production'),
    steps([
        'Open Stock → Record production.',
        'Add a line: choose the product, tap the size, enter the quantity made. '
        'Add more lines for other sizes in the same batch.',
        'Optionally add the tailor’s name, then Save. Stock rises from these '
        'entries — it is never typed in directly.',
    ]),
    h2('Adjust stock'),
    steps([
        'On the Stock screen, find the product/size whose count is wrong and select '
        'Adjust.',
        'Choose Add or Remove, enter the amount and a reason (required), and Save. '
        'The difference is recorded as a movement; the count is not overwritten.',
    ]),
    p('The Seller cannot change prices, manage accounts, or configure sizes and '
      'policies — those belong to the Super Admin.'),

    PB,
    h1('Administration'),
    p('Administration accounts are strictly read-only: they can open and read '
      'everything but change nothing (enforced on the server, not just by hiding '
      'buttons).'),
    h2('What you can do'),
    steps([
        'Open any screen from the sidebar to read it — Open jobs, sales, orders, '
        'alterations, returns, collections.',
        'Open Reports for the daily reconciliation and every other report, and '
        'export to Excel or PDF.',
        'Open the Audit log to see every recorded action.',
    ]),
    p('Attempting to record a sale, change a price or alter any record is refused.'),

    h1('Maintenance (developers)'),
    p('Maintenance accounts have full functional access to help run the shop and '
      'fix problems, and every action is audited like anyone else. Maintenance '
      'also reads and works through the support messages and bug reports.'),

    PB,
    h1('Super Admin'),
    h2('Manage accounts'),
    steps([
        'Open Accounts.',
        'Create an account: enter name, email and role; share the generated '
        'temporary password securely (the user changes it at first login).',
        'Reset password: issue a new temporary password for a locked-out user.',
        'Deactivate: stop a user signing in while keeping their name on their past '
        'work (this is how you retire someone).',
        'Activate: re-enable a deactivated user.',
        'Delete: only for an account with no activity on record; anything with '
        'history must be deactivated instead. You cannot act on your own account.',
    ]),
    h2('Manage the catalogue and prices'),
    steps([
        'Open Catalogue.',
        'New product: enter the garment name and price (no size — size is chosen at '
        'sale). Save.',
        'Change a price by editing the product — the change is audited (old and new '
        'value) and is not retroactive; past sales keep their price.',
        'Archive a product to remove it from the sale screen while keeping it in '
        'past records (products are archived, never deleted).',
    ]),
    h2('Configure the size set'),
    steps([
        'Open Settings → Sizes.',
        'Choose Letters (type a list such as S, M, L, XL) or Numbers (set a range, '
        'e.g. 20 to 46, step 2).',
        'Check the live preview — those are the exact boxes the Seller will see — '
        'and Save.',
    ]),
    h2('Set the return-policy windows'),
    steps([
        'Open Settings.',
        'Set the four windows (exchange/refund × unworn/worn). Save.',
        'Changing a window never reclassifies past returns — each keeps the verdict '
        'it was given.',
    ]),

    PB,
    h1('Reports and exports (all roles)'),
    p('All reports are available to every role. Each export is stamped with the '
      'date, who generated it and the filters, and is recorded.'),
    h2('Produce a report or export'),
    steps([
        'Open Reports.',
        'For the end of day, read the Daily cash reconciliation: net cash that '
        'should be in the box (cash only) is shown separately from the mobile-money '
        'total to check against the phone.',
        'For other reports, pick the report and date range: sales by period, by '
        'garment and size, production, orders and turnaround, returns and '
        'exchanges, cancellations, audit log.',
        'Select Export Excel, or Print / Save as PDF.',
    ]),

    h1('Reporting a problem'),
    steps([
        'On any screen, select Report a problem in the sidebar.',
        'Describe what happened; optionally attach a screenshot. It captures the '
        'page, your account and browser automatically.',
        'Submit — Maintenance and the Super Admin are notified at once. If the app '
        'itself is down, phone the first responder instead.',
    ]),

    PB,
    h1('Quick reference'),
    table(['Prefix', 'Document'], [
        ['SAL', 'Sale'], ['ORD', 'Order (garment to follow)'],
        ['COL', 'Collection of an order'], ['ALT', 'Alteration (deposit slip)'],
        ['RTN', 'Return or exchange'],
    ]),
    p('References are sequential and never reused within their type and year (e.g. '
      'SAL-2026-0001). A cancelled document keeps its number.'),
    bullets([
        'You cannot break the records by making a mistake — corrections are new, '
        'linked entries.',
        'No paper? Search by parent name or phone.',
        'System and shelf disagree? Trust the shelf, confirm the override, and '
        'record production/adjustments so counts catch up.',
    ]),
]

# ============================================================ FRENCH
FR = [
    h1('À propos de ce manuel'),
    p('Ce manuel explique ce que fait le système de vente et de stock d’uniformes '
      'scolaires et, étape par étape, comment chaque utilisateur effectue ses '
      'tâches. Il couvre la Phase 1 — production, ventes, commandes, retouches, '
      'retours et rapports. Lisez la section de votre rôle et les sections '
      'communes (Premiers pas, Règles essentielles, Signaler un problème).'),
    note('L’interface est bilingue. Changez de langue à tout moment depuis la barre '
         'du haut ; chaque écran et chaque reçu existe dans les deux langues. Ce '
         'que vous saisissez — noms, notes, motifs — est conservé tel quel et '
         'jamais traduit.'),

    h1('Le système en bref'),
    p('Les uniformes sont fabriqués et vendus au même endroit. Le système '
      'enregistre la production au fur et à mesure, enregistre les ventes et '
      'imprime les reçus, prend des commandes pour les articles non encore faits, '
      'suit les retouches, gère les retours et échanges, et produit le '
      'rapprochement de caisse de fin de journée.'),
    h2('Règles que chacun doit connaître'),
    bullets([
        'On ne modifie ni ne supprime jamais rien — une erreur se corrige par une '
        'nouvelle écriture liée (une annulation, un retour) ; les deux restent '
        'visibles.',
        'Tout est audité : qui a fait quoi, et quand.',
        'Le paiement est intégral, sur place — espèces ou mobile money, sans '
        'échelonnement.',
        'La taille est choisie au moment de la vente/commande/échange ; le stock '
        'est compté par taille.',
        'Pensé pour le mobile : grands boutons, listes cherchables, peu de saisie.',
    ]),

    PB,
    h1('Premiers pas (tous les utilisateurs)'),
    h2('Se connecter'),
    steps([
        'Ouvrez l’application dans un navigateur. Aucune inscription publique — le '
        'Super Admin crée votre compte.',
        'Saisissez votre e-mail et votre mot de passe, puis Se connecter.',
        'À la première connexion, définissez un nouveau mot de passe (au moins 8 '
        'caractères, avec une lettre et un chiffre) avant de continuer.',
    ]),
    p('Si votre compte est désactivé, vous ne pouvez pas vous connecter — '
      'contactez votre administrateur. Les mots de passe oubliés sont '
      'réinitialisés par le Super Admin (pas de service par e-mail). Vous êtes '
      'déconnecté après inactivité : 12 heures pour le Vendeur, 2 heures pour les '
      'autres.'),
    h2('Se repérer'),
    steps([
        'Utilisez la barre latérale pour naviguer (Vue d’ensemble, Opérations, '
        'Registres, Administration) — vous ne voyez que ce que votre rôle permet.',
        'Utilisez la recherche pour trouver une transaction par numéro de '
        'référence, nom de l’élève, nom du parent ou téléphone.',
        'Ouvrez la cloche pour les notifications ; utilisez les boutons de langue '
        'et de thème en haut.',
    ]),

    PB,
    h1('Les rôles en un coup d’œil'),
    table(['Rôle', 'Qui', 'Ce qu’ils font'], [
        ['Vendeur', 'M. Ateba (1)',
         'Production, ventes, commandes, retouches, retours, échanges, annulations. '
         'Ne peut pas changer les prix.'],
        ['Administration', 'Fondateur, personnel (~5)',
         'Lecture seule partout : vues, rapports, exports, journal d’audit. Aucune '
         'écriture.'],
        ['Maintenance', 'Développeurs (2)', 'Accès fonctionnel complet, entièrement '
         'audité.'],
        ['Super Admin', '1', 'Comptes, catalogue et prix, jeu de tailles, délais de '
         'la politique de retour.'],
    ]),

    PB,
    h1('Le Vendeur'),
    p('Le Vendeur (M. Ateba) tient la boutique. L’écran d’accueil est Travaux en '
      'cours.'),
    h2('Lire la liste des travaux en cours'),
    steps([
        'Ouvrez l’application — Travaux en cours est votre écran d’accueil.',
        'Lisez chaque carte : élève et classe, article et taille, statut, date, et '
        'nombre de jours ouverts. Le plus ancien est en haut.',
        'Touchez une carte pour changer son statut en un geste, ou pour l’ouvrir.',
    ]),
    h2('Enregistrer une vente'),
    steps([
        'Ouvrez Ventes.',
        'Saisissez le nom et la classe de l’élève, et le nom du parent (téléphone '
        'facultatif).',
        'Touchez Ajouter un article, choisissez le produit, puis touchez la taille '
        'sur la barre de tailles (ou saisissez une taille personnalisée) et '
        'indiquez la quantité. Répétez pour d’autres articles.',
        'Pour réduire un prix, saisissez un montant dans Remise avec un motif (le '
        'prix unitaire ne se modifie pas).',
        'Choisissez le mode de paiement — Espèces, MoMo ou Orange Money — et, pour '
        'le mobile money, la référence de la transaction.',
        'Pour les espèces, saisissez éventuellement le montant remis pour voir la '
        'monnaie.',
        'Sélectionnez Enregistrer. Le reçu s’ouvre et le stock de ce produit/taille '
        'diminue.',
    ]),
    note('Vendre plus que ce qu’affiche le système vous avertit mais ne bloque pas '
         'la vente ; confirmez pour continuer et la dérogation est enregistrée.'),
    h2('Imprimer ou réimprimer un reçu'),
    steps([
        'Après l’enregistrement, le reçu s’ouvre — choisissez A5 ou A4 et '
        'Imprimer.',
        'Pour réimprimer plus tard, ouvrez la vente (par la recherche ou la liste) '
        'et sélectionnez Réimprimer. La copie porte la mention DUPLICATA / '
        'DUPLICATE et est enregistrée.',
    ]),
    h2('Prendre une commande (article non disponible)'),
    steps([
        'Ouvrez Commandes et saisissez le client et la ou les lignes avec la '
        'taille, comme pour une vente.',
        'Encaissez le paiement intégral — le reçu porte la mention COMMANDE / '
        'ORDER.',
        'La commande apparaît dans Travaux en cours ; faites-la avancer Commandée → '
        'En production → Prête.',
        'À la venue du parent, ouvrez la commande et enregistrez le retrait : un '
        'bon de retrait s’imprime (indiquant le COL et le ORD qu’il clôt) et le '
        'stock diminue à ce moment.',
    ]),
    h2('Réceptionner une retouche'),
    steps([
        'Ouvrez Retouches → Nouvelle.',
        'Enregistrez le parent/l’élève, l’article et ce qui doit être fait '
        '(obligatoire).',
        'Indiquez s’il y a des frais ; si oui, encaissez et indiquez le mode.',
        'Remettez au parent le bon de dépôt imprimé.',
        'Faites avancer Reçue → En cours → Prête, et enregistrez Rendue au retrait. '
        'Les retouches ne touchent jamais au stock.',
    ]),
    h2('Traiter un retour ou un échange'),
    steps([
        'Ouvrez la vente d’origine et sélectionnez Démarrer un retour (ou Retours → '
        'Nouveau).',
        'Lisez le verdict à l’écran — depuis combien de temps la vente a eu lieu et '
        'si l’échange/remboursement est dans les délais.',
        'Choisissez le ou les articles rendus et déclarez l’état (Non porté ou '
        'Porté).',
        'Pour un échange, choisissez le produit et la taille de remplacement ; '
        'encaissez la différence si plus cher, remboursez-la si moins cher. Pour un '
        'retour, effectuez le remboursement et indiquez son mode.',
        'Saisissez un motif (obligatoire). Hors délais, ajoutez le motif demandé et '
        'continuez — la dérogation est signalée dans les rapports.',
        'Enregistrez — un reçu de retour/échange s’imprime et le stock est mis à '
        'jour dans les deux sens.',
    ]),
    h2('Annuler une vente'),
    steps([
        'Ouvrez la vente et sélectionnez Annuler.',
        'Saisissez un motif (obligatoire) et confirmez.',
        'Le stock est restitué ; la vente reste visible avec la mention Annulée, '
        'garde sa référence et est exclue du chiffre d’affaires.',
    ]),
    h2('Enregistrer la production'),
    steps([
        'Ouvrez Stock → Enregistrer la production.',
        'Ajoutez une ligne : choisissez le produit, touchez la taille, saisissez la '
        'quantité fabriquée. Ajoutez d’autres lignes pour d’autres tailles dans le '
        'même lot.',
        'Ajoutez éventuellement le nom du tailleur, puis Enregistrez. Le stock '
        'monte à partir de ces écritures — il ne se saisit jamais directement.',
    ]),
    h2('Ajuster le stock'),
    steps([
        'Sur l’écran Stock, trouvez le produit/taille dont le compte est faux et '
        'sélectionnez Ajuster.',
        'Choisissez Ajouter ou Retirer, saisissez la quantité et un motif '
        '(obligatoire), puis Enregistrez. La différence est enregistrée comme '
        'mouvement ; le compte n’est pas écrasé.',
    ]),
    p('Le Vendeur ne peut pas changer les prix, gérer les comptes, ni configurer '
      'les tailles et politiques — cela revient au Super Admin.'),

    PB,
    h1('Administration'),
    p('Les comptes Administration sont en lecture seule stricte : ils peuvent tout '
      'ouvrir et lire mais ne rien changer (imposé côté serveur, pas seulement en '
      'masquant des boutons).'),
    h2('Ce que vous pouvez faire'),
    steps([
        'Ouvrez n’importe quel écran depuis la barre latérale pour le consulter — '
        'Travaux en cours, ventes, commandes, retouches, retours, retraits.',
        'Ouvrez Rapports pour le rapprochement quotidien et tous les autres '
        'rapports, et exportez en Excel ou PDF.',
        'Ouvrez le Journal d’audit pour voir chaque action enregistrée.',
    ]),
    p('Toute tentative d’enregistrer une vente, de changer un prix ou de modifier '
      'un enregistrement est refusée.'),

    h1('Maintenance (développeurs)'),
    p('Les comptes Maintenance ont un accès fonctionnel complet pour aider à faire '
      'tourner la boutique et corriger les problèmes, et chaque action est audité '
      'comme pour tout le monde. La Maintenance lit aussi et traite les messages '
      'd’assistance et les rapports de bug.'),

    PB,
    h1('Super Admin'),
    h2('Gérer les comptes'),
    steps([
        'Ouvrez Comptes.',
        'Créer un compte : saisissez le nom, l’e-mail et le rôle ; partagez le mot '
        'de passe temporaire généré de façon sécurisée (l’utilisateur le change à '
        'la première connexion).',
        'Réinitialiser le mot de passe : émettez un nouveau mot de passe temporaire.',
        'Désactiver : empêchez un utilisateur de se connecter tout en gardant son '
        'nom sur ses travaux passés (c’est ainsi qu’on retire quelqu’un).',
        'Activer : réactivez un utilisateur désactivé.',
        'Supprimer : seulement pour un compte sans activité enregistrée ; tout ce '
        'qui a un historique doit être désactivé. Vous ne pouvez pas agir sur '
        'votre propre compte.',
    ]),
    h2('Gérer le catalogue et les prix'),
    steps([
        'Ouvrez Catalogue.',
        'Nouveau produit : saisissez le nom de l’article et le prix (pas de taille '
        '— elle est choisie à la vente). Enregistrez.',
        'Changez un prix en modifiant le produit — le changement est audité (ancien '
        'et nouveau) et non rétroactif ; les ventes passées gardent leur prix.',
        'Archivez un produit pour le retirer de l’écran de vente tout en le gardant '
        'dans les registres (les produits sont archivés, jamais supprimés).',
    ]),
    h2('Configurer le jeu de tailles'),
    steps([
        'Ouvrez Paramètres → Tailles.',
        'Choisissez Lettres (saisissez une liste comme S, M, L, XL) ou Chiffres '
        '(définissez une plage, p. ex. 20 à 46, pas de 2).',
        'Vérifiez l’aperçu en direct — ce sont exactement les cases que verra le '
        'Vendeur — et Enregistrez.',
    ]),
    h2('Définir les délais de retour'),
    steps([
        'Ouvrez Paramètres.',
        'Définissez les quatre délais (échange/remboursement × non porté/porté). '
        'Enregistrez.',
        'Modifier un délai ne reclasse jamais les retours passés — chacun garde son '
        'verdict.',
    ]),

    PB,
    h1('Rapports et exports (tous les rôles)'),
    p('Tous les rapports sont accessibles à chaque rôle. Chaque export porte la '
      'date, l’auteur et les filtres, et est enregistré.'),
    h2('Produire un rapport ou un export'),
    steps([
        'Ouvrez Rapports.',
        'Pour la fin de journée, lisez le Rapprochement de caisse : l’argent '
        'liquide qui doit être en caisse (espèces uniquement) est affiché à part du '
        'total mobile money à vérifier sur le téléphone.',
        'Pour les autres rapports, choisissez le rapport et la période : ventes par '
        'période, par article et taille, production, commandes et délais, retours '
        'et échanges, annulations, journal d’audit.',
        'Sélectionnez Exporter Excel, ou Imprimer / Enregistrer en PDF.',
    ]),

    h1('Signaler un problème'),
    steps([
        'Sur n’importe quel écran, sélectionnez Signaler un problème dans la barre '
        'latérale.',
        'Décrivez ce qui s’est passé ; joignez éventuellement une capture. Le '
        'système capte la page, votre compte et le navigateur automatiquement.',
        'Envoyez — la Maintenance et le Super Admin sont notifiés aussitôt. Si '
        'l’application est hors service, appelez le premier intervenant.',
    ]),

    PB,
    h1('Référence rapide'),
    table(['Préfixe', 'Document'], [
        ['SAL', 'Vente'], ['ORD', 'Commande (article à suivre)'],
        ['COL', 'Retrait d’une commande'], ['ALT', 'Retouche (bon de dépôt)'],
        ['RTN', 'Retour ou échange'],
    ]),
    p('Les références sont séquentielles et jamais réutilisées dans leur type et '
      'leur année (p. ex. SAL-2026-0001). Un document annulé garde son numéro.'),
    bullets([
        'Vous ne pouvez pas casser les registres en faisant une erreur — les '
        'corrections sont de nouvelles écritures liées.',
        'Pas de papier ? Cherchez par nom du parent ou téléphone.',
        'Le système et l’étagère divergent ? Faites confiance à l’étagère, '
        'confirmez la dérogation, et enregistrez production/ajustements pour que '
        'les comptes se rattrapent.',
    ]),
]

COVERS = {
    'en': ('User Manual', 'School Uniform Sales & Stock — Phase 1',
           'How each user performs their tasks · English (Français ci-dessous)'),
    'fr': ('Manuel d’utilisation', 'Vente et stock d’uniformes scolaires — Phase 1',
           'Comment chaque utilisateur effectue ses tâches · Français'),
}

# ---------------------------------------------------------------- markdown
def md_blocks(out, blocks):
    for b in blocks:
        k = b[0]
        if k == 'pagebreak': out.append('\n---\n')
        elif k == 'h1': out.append(f'\n## {b[1]}\n')
        elif k == 'h2': out.append(f'\n### {b[1]}\n')
        elif k == 'p': out.append(b[1] + '\n')
        elif k == 'note': out.append(f'> **Note:** {b[1]}\n')
        elif k == 'bullets':
            out.extend(f'- {i}' for i in b[1]); out.append('')
        elif k == 'steps':
            out.extend(f'{n}. {i}' for n, i in enumerate(b[1], 1)); out.append('')
        elif k == 'table':
            headers, rows = b[1], b[2]
            out.append('| ' + ' | '.join(headers) + ' |')
            out.append('| ' + ' | '.join('---' for _ in headers) + ' |')
            for r in rows:
                out.append('| ' + ' | '.join(c.replace('\n', ' ') for c in r) + ' |')
            out.append('')

def render_md(path):
    out = [f'# {SCHOOL}', '', f'## {COVERS["en"][0]} — {COVERS["en"][1]}', '',
           f'_{COVERS["en"][2]}_', '']
    md_blocks(out, EN)
    out.append('\n---\n')
    out.append(f'# {SCHOOL}')
    out.append('')
    out.append(f'## {COVERS["fr"][0]} — {COVERS["fr"][1]}')
    out.append('')
    out.append(f'_{COVERS["fr"][2]}_')
    out.append('')
    md_blocks(out, FR)
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out) + '\n')

# ---------------------------------------------------------------- docx
def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexfill)
    tcPr.append(shd)

def docx_cover(doc, lang):
    title, sub, tag = COVERS[lang]
    if os.path.exists(LOGO):
        pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(LOGO, width=Inches(1.6))
    for text, size, color, bold, italic in [
        (SCHOOL, 22, GREEN_DARK, True, False),
        (title, 16, GREEN, True, False),
        (sub, 12, GREY, False, False),
        (tag, 10, GREY, False, True),
    ]:
        pr = doc.add_paragraph(); pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pr.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color
        r.bold = bold; r.italic = italic

def docx_heading(doc, text, size, color, before):
    pr = doc.add_paragraph()
    pr.paragraph_format.space_before = Pt(before)
    pr.paragraph_format.space_after = Pt(4)
    r = pr.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color

def docx_blocks(doc, blocks):
    for b in blocks:
        k = b[0]
        if k == 'pagebreak': doc.add_page_break()
        elif k == 'h1': docx_heading(doc, b[1], 16, GREEN_DARK, 16)
        elif k == 'h2': docx_heading(doc, b[1], 13, GREEN, 10)
        elif k == 'p': doc.add_paragraph(b[1])
        elif k == 'note':
            pr = doc.add_paragraph()
            run = pr.add_run('Note: '); run.bold = True; run.font.color.rgb = GREEN
            pr.add_run(b[1]).italic = True
        elif k == 'bullets':
            for i in b[1]: doc.add_paragraph(i, style='List Bullet')
        elif k == 'steps':
            for i in b[1]: doc.add_paragraph(i, style='List Number')
        elif k == 'table':
            headers, rows = b[1], b[2]
            t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for j, htext in enumerate(headers):
                c = t.rows[0].cells[j]; c.paragraphs[0].text = ''
                run = c.paragraphs[0].add_run(htext); run.bold = True
                run.font.color.rgb = WHITE; shade(c, GREEN_HEX)
            for row in rows:
                cells = t.add_row().cells
                for j, val in enumerate(row): cells[j].text = val

def render_docx(path):
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    docx_cover(doc, 'en'); doc.add_page_break()
    docx_blocks(doc, EN)
    doc.add_page_break(); docx_cover(doc, 'fr'); doc.add_page_break()
    docx_blocks(doc, FR)
    doc.save(path)

if __name__ == '__main__':
    md = os.path.join(ROOT, 'docs', 'user-manual.md')
    dx = os.path.join(ROOT, 'docs', 'user-manual.docx')
    render_md(md); render_docx(dx)
    print('Wrote', md); print('Wrote', dx)
